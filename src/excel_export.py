from __future__ import annotations

from datetime import datetime
from io import BytesIO
import re

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


KOREAN_COLUMNS = {
    "company_id": "업체 ID", "permit_number": "인허가번호", "company_name": "업체명",
    "company_name_normalized": "업체명 정규화", "address": "주소", "address_normalized": "주소 정규화",
    "district": "지역", "industry": "업종", "entity_type": "업체구분", "is_2026_target": "2026 점검대상 여부",
    "planned_inspection_type": "예정 점검유형", "manager_alias": "담당자(가명)", "air_scale": "대기 규모",
    "water_scale": "폐수 규모", "air_grade": "대기 등급", "water_grade": "폐수 등급",
    "target_media_waste": "폐기물 대상", "source_tags": "원본 태그", "source_record_count": "원본 건수",
    "major_complaint_note": "주요 민원 참고", "manual_inspection_note": "수기 점검 참고",
    "inspection_id": "지도점검 ID", "inspection_date": "점검일", "inspection_date_raw": "점검일 원문",
    "inspection_type": "점검유형", "inspection_media": "점검매체", "inspection_result_status": "점검결과",
    "inspection_result_detail": "점검결과 상세", "key_findings": "주요 점검내용",
    "suspected_violation": "위반 또는 지적사항", "on_site_action": "현장조치", "follow_up_action": "후속조치",
    "inspector_alias": "점검자(가명)", "review_status": "검토상태", "source_file": "원본 파일",
    "source_sheet": "원본 시트", "source_row": "원본 행", "match_method": "연결 방식",
    "match_score": "연결 점수", "inspection_purpose": "점검목적", "correction_due_date": "시정기한",
    "next_check_points": "다음 점검 참고사항", "data_source": "데이터 출처", "reviewed_at": "검토일시",
    "disposition_id": "행정처분 ID", "disposition_date": "처분일", "violation_content": "위반내용",
    "legal_basis": "관련 법령", "administrative_disposition": "행정처분 내용", "accusation": "고발내용",
    "penalty": "과태료 내용", "violation_category": "위반유형", "representative_alias": "대표자(가명)",
    "source_media": "원본 매체", "disposition_status": "처분상태", "note": "비고",
    "alias_id": "별칭 ID", "alias_type": "별칭 유형", "alias_value": "별칭", "normalized_value": "정규화 별칭",
    "source": "별칭 출처", "is_primary": "대표 별칭 여부", "created_at": "생성일시", "updated_at": "수정일시",
}

DATE_COLUMNS = {
    "inspection_date", "correction_due_date", "disposition_date", "reviewed_at", "created_at", "updated_at",
}
INVALID_FILENAME = re.compile(r'[\\/:*?"<>|]')


def _clean_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in result.columns:
        if column in DATE_COLUMNS:
            parsed = pd.to_datetime(result[column], errors="coerce")
            formatted = parsed.dt.strftime("%Y-%m-%d")
            result[column] = formatted.where(parsed.notna(), result[column])
    result = result.where(pd.notna(result), "")
    return result.rename(columns={column: KOREAN_COLUMNS.get(column, column) for column in result.columns})


def _write_workbook(sheets: list[tuple[str, pd.DataFrame]]) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, frame in sheets:
            cleaned = _clean_frame(frame)
            cleaned.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.book[sheet_name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = f"A1:{get_column_letter(max(1, worksheet.max_column))}{max(1, worksheet.max_row)}"
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for column_cells in worksheet.iter_cols(min_row=1, max_row=worksheet.max_row):
                width = max((len(str(cell.value or "")) for cell in column_cells), default=10)
                worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(width + 2, 10), 45)
    return buffer.getvalue()


def _safe_company_name(company_name: object) -> str:
    value = INVALID_FILENAME.sub("_", str(company_name or "").strip())
    return value or "업체"


def company_history_filename(company_name: object, now: datetime | None = None) -> str:
    return f"{_safe_company_name(company_name)}_통합이력_{(now or datetime.now()).strftime('%Y%m%d')}.xlsx"


def all_data_filename(now: datetime | None = None) -> str:
    return f"전체데이터_{(now or datetime.now()).strftime('%Y%m%d_%H%M')}.xlsx"


def build_company_history_excel(
    company: pd.Series | dict[str, object], inspections: pd.DataFrame, dispositions: pd.DataFrame
) -> bytes:
    return _write_workbook([
        ("업체기본정보", pd.DataFrame([dict(company)])),
        ("지도점검이력", inspections),
        ("행정처분이력", dispositions),
    ])


def build_search_results_excel(results: pd.DataFrame, query: str, filters: dict[str, str]) -> bytes:
    conditions = pd.DataFrame([
        ("검색어", query or ""),
        *[(KOREAN_COLUMNS.get(field, field), value or "전체") for field, value in filters.items()],
    ], columns=["검색조건", "적용값"])
    return _write_workbook([("검색결과", results), ("검색조건", conditions)])


def build_all_data_excel(data: dict[str, pd.DataFrame]) -> bytes:
    return _write_workbook([
        ("업체마스터", data["companies"]),
        ("지도점검이력", data["inspections"]),
        ("행정처분이력", data["dispositions"]),
        ("업체별칭", data["aliases"]),
    ])

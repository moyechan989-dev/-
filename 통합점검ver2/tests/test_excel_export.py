from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.excel_export import (
    all_data_filename,
    build_all_data_excel,
    build_company_history_excel,
    build_search_results_excel,
    company_history_filename,
)


def workbook(content: bytes):
    return load_workbook(BytesIO(content))


def test_company_history_excel_preserves_history_counts_and_formatting():
    company = pd.Series({"company_id": "COM-1", "company_name": "가나다 업체", "address": None})
    inspections = pd.DataFrame([{"inspection_id": "INS-1", "company_id": "COM-1", "inspection_date": "2026-08-27"}])
    dispositions = pd.DataFrame([
        {"disposition_id": "DSP-1", "company_id": "COM-1", "disposition_date": "2026-08-26"},
        {"disposition_id": "DSP-2", "company_id": "COM-1", "disposition_date": "2026-08-25"},
    ])
    book = workbook(build_company_history_excel(company, inspections, dispositions))
    assert book.sheetnames == ["업체기본정보", "지도점검이력", "행정처분이력"]
    assert book["지도점검이력"].max_row - 1 == len(inspections)
    assert book["행정처분이력"].max_row - 1 == len(dispositions)
    assert book["업체기본정보"]["C2"].value in (None, "")
    assert book["지도점검이력"].freeze_panes == "A2"
    assert book["지도점검이력"].auto_filter.ref == "A1:C2"


def test_company_history_excel_supports_no_history():
    company = pd.Series({"company_id": "COM-1", "company_name": "이력 없는 업체"})
    inspections = pd.DataFrame(columns=["inspection_id", "company_id", "inspection_date"])
    dispositions = pd.DataFrame(columns=["disposition_id", "company_id", "disposition_date"])
    book = workbook(build_company_history_excel(company, inspections, dispositions))
    assert book["지도점검이력"].max_row == 1
    assert book["행정처분이력"].max_row == 1


def test_search_results_excel_preserves_result_count_and_conditions():
    results = pd.DataFrame([
        {"company_id": "COM-1", "company_name": "업체 1"},
        {"company_id": "COM-2", "company_name": "업체 2"},
    ])
    book = workbook(build_search_results_excel(results, "업체", {"district": "가동", "industry": ""}))
    assert book["검색결과"].max_row - 1 == len(results)
    assert book["검색조건"]["A1"].value == "검색조건"
    assert book["검색조건"]["B2"].value == "업체"
    assert book["검색조건"]["B4"].value == "전체"


def test_all_data_excel_preserves_each_table_count():
    data = {
        "companies": pd.DataFrame([{"company_id": "COM-1"}, {"company_id": "COM-2"}]),
        "inspections": pd.DataFrame([{"inspection_id": "INS-1"}]),
        "dispositions": pd.DataFrame([{"disposition_id": "DSP-1"}, {"disposition_id": "DSP-2"}]),
        "aliases": pd.DataFrame([{"alias_id": 1}, {"alias_id": 2}, {"alias_id": 3}]),
    }
    book = workbook(build_all_data_excel(data))
    assert book.sheetnames == ["업체마스터", "지도점검이력", "행정처분이력", "업체별칭"]
    assert [book[name].max_row - 1 for name in book.sheetnames] == [2, 1, 2, 3]


def test_korean_filenames_are_safe_and_timestamped():
    now = datetime(2026, 8, 27, 13, 30)
    assert company_history_filename("가나다/업체", now) == "가나다_업체_통합이력_20260827.xlsx"
    assert all_data_filename(now) == "전체데이터_20260827_1330.xlsx"
